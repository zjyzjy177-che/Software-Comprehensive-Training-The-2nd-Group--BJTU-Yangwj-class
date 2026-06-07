"""
仿真报告导出模块 - export_report.py
导出 Excel(.xlsx) / Word(.docx)
依赖：pip install openpyxl python-docx
字体：中文=微软雅黑 英文/数字=Courier New Bold
"""
import os
from datetime import datetime
from typing import Dict, Any, List

# 字体配置
CN_FONT = "Microsoft YaHei"
EN_FONT = "Courier New"


def _safe_stat(stats, key, default="N/A"):
    val = stats.get(key, default)
    if isinstance(val, float):
        return round(val, 2)
    return val


def _make_fonts():
    """创建字体对象（需要先 import openpyxl.styles）"""
    from openpyxl.styles import Font
    return {
        'title': Font(name=CN_FONT, size=14, bold=True, color="1A5276"),
        'header': Font(name=CN_FONT, size=11, bold=True, color="FFFFFF"),
        'label': Font(name=CN_FONT, size=11, bold=True),
        'value': Font(name=EN_FONT, size=11, bold=True),
        'text': Font(name=CN_FONT, size=10),
        'small': Font(name=CN_FONT, size=9, color="666666"),
    }


def export_excel(engine_or_stats, filepath: str = "simulation_report.xlsx",
                 tick_history: List = None) -> str:
    """导出 Excel 报告"""
    try:
        import openpyxl
        from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
        from openpyxl.chart import BarChart, Reference
        from openpyxl.utils import get_column_letter
    except ImportError:
        raise ImportError("请先安装 openpyxl: pip install openpyxl")

    if hasattr(engine_or_stats, 'get_statistics'):
        stats = engine_or_stats.get_statistics()
        th = tick_history or getattr(engine_or_stats, 'tick_history', [])
    else:
        stats = engine_or_stats
        th = tick_history or []

    f = _make_fonts()
    blue_fill = PatternFill(start_color="3498DB", end_color="3498DB", fill_type="solid")
    green_fill = PatternFill(start_color="27AE60", end_color="27AE60", fill_type="solid")
    light_fill = PatternFill(start_color="EBF5FB", end_color="EBF5FB", fill_type="solid")
    thin_border = Border(
        left=Side(style='thin'), right=Side(style='thin'),
        top=Side(style='thin'), bottom=Side(style='thin'))
    center = Alignment(horizontal='center', vertical='center')
    left_align = Alignment(horizontal='left', vertical='center')

    wb = openpyxl.Workbook()

    # ==================== Sheet 1: 统计摘要 ====================
    ws = wb.active
    ws.title = "仿真统计"
    ws.sheet_properties.tabColor = "3498DB"

    # 标题
    ws.merge_cells('A1:D1')
    c = ws['A1']
    c.value = "BJTU 食堂就餐流量仿真 — 统计报告"
    c.font = f['title']
    c.alignment = Alignment(horizontal='center')
    ws.row_dimensions[1].height = 30

    ws.merge_cells('A2:D2')
    ws['A2'].value = f"生成时间: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}"
    ws['A2'].font = f['small']
    ws['A2'].alignment = Alignment(horizontal='center')

    # 核心指标
    metrics = [
        ("仿真周期数", "current_tick", ""),
        ("总生成学生数", "total_students_generated", "人"),
        ("总服务学生数", "total_students_served", "人"),
        ("最大排队人数", "max_queue_length", "人"),
        ("平均等待时间", "average_wait_time", "tick"),
        ("最终活跃学生数", "final_active_students", "人"),
    ]
    row = 4
    ws.cell(row=row, column=1, value="核心指标").font = Font(name=CN_FONT, size=12, bold=True, color="1A5276")
    row = 5
    for col, h in enumerate(["指标", "数值", "单位"], 1):
        c = ws.cell(row=row, column=col, value=h)
        c.font = f['header']; c.fill = blue_fill; c.alignment = center; c.border = thin_border
    for i, (label, key, unit) in enumerate(metrics):
        row = 6 + i
        ws.cell(row=row, column=1, value=label).font = f['label']
        ws.cell(row=row, column=2, value=_safe_stat(stats, key)).font = f['value']
        ws.cell(row=row, column=3, value=unit).font = f['text']
        if i % 2 == 0:
            for c in range(1, 4):
                ws.cell(row=row, column=c).fill = light_fill
        for c in range(1, 4):
            ws.cell(row=row, column=c).border = thin_border

    # 食堂详表
    cs = stats.get('canteen_statistics', {})
    if cs:
        row = 14
        ws.cell(row=row, column=1, value="食堂运营详情").font = Font(name=CN_FONT, size=12, bold=True, color="1A5276")
        row = 15
        headers = ["食堂", "服务人数", "当前排队", "最大排队", "容量", "利用率"]
        for col, h in enumerate(headers, 1):
            c = ws.cell(row=row, column=col, value=h)
            c.font = f['header']; c.fill = green_fill; c.alignment = center; c.border = thin_border
        for i, (cname, cinfo) in enumerate(cs.items()):
            row = 16 + i
            cap = cinfo.get('capacity', 1)
            util = f"{cinfo.get('max_queue', 0) / max(cap, 1) * 100:.0f}%"
            vals = [cname, cinfo.get('served', 0), cinfo.get('current_queue', 0),
                    cinfo.get('max_queue', 0), cap, util]
            for col, v in enumerate(vals, 1):
                ws.cell(row=row, column=col, value=v).font = f['text']
                ws.cell(row=row, column=col).border = thin_border
                ws.cell(row=row, column=col).alignment = center

    # 列宽
    ws.column_dimensions['A'].width = 20
    ws.column_dimensions['B'].width = 16
    ws.column_dimensions['C'].width = 14
    ws.column_dimensions['D'].width = 14
    ws.column_dimensions['E'].width = 12
    ws.column_dimensions['F'].width = 12

    # ==================== Sheet 2: 排队曲线数据 ====================
    if th:
        ws2 = wb.create_sheet("排队曲线")
        ws2.sheet_properties.tabColor = "E74C3C"
        ws2.cell(row=1, column=1, value="Tick").font = f['header']
        ws2.cell(row=1, column=1).fill = blue_fill
        ws2.cell(row=1, column=2, value="排队人数").font = f['header']
        ws2.cell(row=1, column=2).fill = blue_fill
        ws2.cell(row=1, column=3, value="活跃学生").font = f['header']
        ws2.cell(row=1, column=3).fill = blue_fill
        for i, td in enumerate(th):
            ws2.cell(row=i+2, column=1, value=i+1).font = f['value']
            ws2.cell(row=i+2, column=2, value=td.get('total_queue_length', 0)).font = f['text']
            ws2.cell(row=i+2, column=3, value=td.get('active_students', 0)).font = f['text']
        ws2.column_dimensions['A'].width = 10
        ws2.column_dimensions['B'].width = 14
        ws2.column_dimensions['C'].width = 14

    # ==================== Sheet 3: 错峰建议 ====================
    ws3 = wb.create_sheet("错峰建议")
    ws3.sheet_properties.tabColor = "F39C12"
    ws3.merge_cells('A1:C1')
    ws3['A1'].value = "错峰下课方案建议"
    ws3['A1'].font = f['title']
    ws3['A1'].alignment = Alignment(horizontal='center')
    tips = [
        "建议采用 15-30 分钟的下课时间差方案",
        "合理错峰可将排队峰值降低 40%-65%",
        "思源楼/逸夫楼等大教学楼优先错峰效果最为显著",
        "错峰对比详细数据请运行: python3 peak_shift.py",
        "可生成四图对比（峰值/等待/曲线/服务）",
    ]
    for i, tip in enumerate(tips, 3):
        ws3.cell(row=i, column=1, value=f"• {tip}").font = f['text']
    ws3.column_dimensions['A'].width = 55

    wb.save(filepath)
    return os.path.abspath(filepath)


def export_docx(engine_or_stats, filepath: str = "simulation_report.docx",
                tick_history: List = None) -> str:
    """导出 Word 报告"""
    try:
        from docx import Document
        from docx.shared import Inches, Pt, RGBColor, Cm
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        from docx.oxml.ns import qn
    except ImportError:
        raise ImportError("请先安装 python-docx: pip install python-docx")

    if hasattr(engine_or_stats, 'get_statistics'):
        stats = engine_or_stats.get_statistics()
        th = tick_history or getattr(engine_or_stats, 'tick_history', [])
    else:
        stats = engine_or_stats
        th = tick_history or []

    doc = Document()

    # 全局默认字体
    style = doc.styles['Normal']
    style.font.name = CN_FONT
    style.font.size = Pt(11)
    style.element.rPr.rFonts.set(qn('w:eastAsia'), CN_FONT)

    # 封面标题
    title = doc.add_heading("BJTU 食堂就餐流量仿真报告", level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for run in title.runs:
        run.font.name = CN_FONT
        run.element.rPr.rFonts.set(qn('w:eastAsia'), CN_FONT)
        run.font.size = Pt(22)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(f"生成时间：{datetime.now().strftime('%Y-%m-%d %H:%M:%S')}")
    run.font.size = Pt(10); run.font.color.rgb = RGBColor(0x66, 0x66, 0x66)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("北京交通大学 — 软件综合实训第二小组")
    run.font.size = Pt(10)

    doc.add_paragraph()

    # 一、核心统计
    doc.add_heading("一、核心仿真统计", level=1)
    table = doc.add_table(rows=7, cols=2, style='Light Shading Accent 1')
    metrics = [
        ("仿真周期数", "current_tick"),
        ("总生成学生数", "total_students_generated"),
        ("总服务学生数", "total_students_served"),
        ("最大排队人数", "max_queue_length"),
        ("平均等待时间 (tick)", "average_wait_time"),
        ("最终活跃学生数", "final_active_students"),
    ]
    for i, (label, key) in enumerate(metrics):
        table.rows[i].cells[0].text = label
        val = str(_safe_stat(stats, key))
        table.rows[i].cells[1].text = val
        # 数字列用 Courier New
        for p in table.rows[i].cells[1].paragraphs:
            for run in p.runs:
                run.font.name = EN_FONT
                run.font.bold = True

    # 二、食堂运营
    cs = stats.get('canteen_statistics', {})
    if cs:
        doc.add_heading("二、食堂运营详情", level=1)
        for cname, cinfo in cs.items():
            cap = cinfo.get('capacity', 1)
            util = cinfo.get('max_queue', 0) / max(cap, 1) * 100
            doc.add_paragraph(
                f"▸ {cname}：已服务 {cinfo.get('served', 0)} 人，"
                f"当前排队 {cinfo.get('current_queue', 0)} 人，"
                f"历史峰值 {cinfo.get('max_queue', 0)} 人，"
                f"容量 {cap} 人，峰值利用率 {util:.1f}%",
                style='List Bullet')

    # 三、错峰建议
    doc.add_heading("三、错峰下课方案建议", level=1)
    doc.add_paragraph(
        "根据错峰对比仿真分析，建议采用 15-30 分钟的下课时间差。合理错峰可将排队峰值降低 40-65%，显著缓解食堂高峰压力。思源楼、逸夫楼等大型教学楼优先安排错峰效果最为显著。")
    doc.add_paragraph("详细数据请运行 python3 peak_shift.py 查看四图对比分析。")

    # 页脚
    doc.add_paragraph()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("— BJTU 软件综合实训第二小组 —")
    run.font.size = Pt(9); run.font.color.rgb = RGBColor(0x99, 0x99, 0x99)

    doc.save(filepath)
    return os.path.abspath(filepath)


def export_pdf(engine_or_stats, filepath: str = "simulation_report.pdf",
               tick_history: List = None) -> str:
    """导出 PDF 报告（使用 matplotlib 渲染）"""
    import matplotlib
    matplotlib.use('Agg')
    import matplotlib.pyplot as plt
    from matplotlib.backends.backend_pdf import PdfPages

    if hasattr(engine_or_stats, 'get_statistics'):
        stats = engine_or_stats.get_statistics()
        th = tick_history or getattr(engine_or_stats, 'tick_history', [])
    else:
        stats = engine_or_stats
        th = tick_history or []

    try:
        plt.rcParams['font.sans-serif'] = ['PingFang SC', 'Microsoft YaHei', 'Heiti SC', 'SimHei']
        plt.rcParams['axes.unicode_minus'] = False
    except Exception:
        pass

    with PdfPages(filepath) as pdf:
        # 封面
        fig = plt.figure(figsize=(8.27, 11.69))  # A4
        fig.suptitle("BJTU 食堂就餐流量仿真报告", fontsize=20, fontweight='bold', y=0.95)
        fig.text(0.5, 0.85, f"生成时间: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}",
                ha='center', fontsize=11, color='#666')
        fig.text(0.5, 0.80, "北京交通大学 — 软件综合实训第二小组",
                ha='center', fontsize=13)

        # 核心指标表
        metrics = [
            ("仿真周期数", "current_tick", ""),
            ("总生成学生数", "total_students_generated", "人"),
            ("总服务学生数", "total_students_served", "人"),
            ("最大排队人数", "max_queue_length", "人"),
            ("平均等待时间", "average_wait_time", "tick"),
            ("最终活跃学生数", "final_active_students", "人"),
        ]
        y = 0.68
        fig.text(0.1, y + 0.04, "核心仿真指标", fontsize=14, fontweight='bold')
        y -= 0.04
        for i, (label, key, unit) in enumerate(metrics):
            val = str(_safe_stat(stats, key))
            fig.text(0.12, y - i * 0.05, f"{label}:  {val} {unit}", fontsize=12)
        pdf.savefig(fig)
        plt.close(fig)

        # 排队曲线图
        if th:
            fig2, ax = plt.subplots(figsize=(10, 5))
            ticks = list(range(1, len(th) + 1))
            queues = [td.get('total_queue_length', 0) for td in th]
            ax.fill_between(ticks, 0, queues, alpha=0.3, color='#E74C3C')
            ax.plot(ticks, queues, color='#E74C3C', linewidth=2)
            ax.set_title("排队人数变化曲线", fontsize=14, fontweight='bold')
            ax.set_xlabel("仿真周期 (tick)"); ax.set_ylabel("排队人数")
            ax.grid(True, alpha=0.3)
            pdf.savefig(fig2)
            plt.close(fig2)

    return os.path.abspath(filepath)


def export_both(engine, tick_history=None, output_dir: str = ".") -> Dict[str, str]:
    """同时导出 Excel 和 Word"""
    ts = datetime.now().strftime("%Y%m%d_%H%M%S")
    th = tick_history or getattr(engine, 'tick_history', [])
    return {
        'xlsx': export_excel(engine, os.path.join(output_dir, f"仿真报告_{ts}.xlsx"), th),
        'docx': export_docx(engine, os.path.join(output_dir, f"仿真报告_{ts}.docx"), th),
    }
